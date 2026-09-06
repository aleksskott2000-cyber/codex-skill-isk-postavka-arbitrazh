#!/usr/bin/env python3
"""Deterministic formatting and residue checks for a generated Russian claim DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document


FONT = "Times New Roman"
ROLE_PATTERN = re.compile(
    r"\b(?:Ист(?:ец|ца|цу|цом|це)|Ответчик(?:а|у|ом|е)?|Договор(?:а|у|ом|е)?|"
    r"Покупател(?:ь|я|ю|ем|е)|Продав(?:ец|ца|цу|цом|це)|Поставщик(?:а|у|ом|е)?)\b"
)
PLACEHOLDER_PATTERN = re.compile(
    r"\[[^\]]*(?:уточнить|вставить|заполнить|ФИО|дата|сумма|номер)[^\]]*\]|"
    r"_{3,}|<\s*(?:уточнить|вставить|заполнить|ФИО|дата|сумма|номер)[^>]*>",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        seen_cells = set()
        for row in table.rows:
            for cell in row.cells:
                cell_id = id(cell._tc)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def main_body_paragraphs(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().lower().startswith("представитель"):
            break
        yield paragraph


def inherited_value(run, paragraph, attribute):
    candidates = [
        run.font,
        run.style.font if run.style is not None else None,
        paragraph.style.font if paragraph.style is not None else None,
    ]
    for font in candidates:
        if font is None:
            continue
        value = getattr(font, attribute)
        if value is not None:
            return value
    return None


def effective_font(run, paragraph):
    return inherited_value(run, paragraph, "name")


def effective_size(run, paragraph):
    value = inherited_value(run, paragraph, "size")
    return value.pt if value is not None else None


def effective_italic(run, paragraph):
    return bool(inherited_value(run, paragraph, "italic"))


def audit(path: Path, template: Path | None) -> dict:
    doc = Document(path)
    errors = []
    warnings = []

    text = "\n".join(p.text for p in all_paragraphs(doc))
    uppercase_roles = sorted(set(ROLE_PATTERN.findall(text)))
    if uppercase_roles:
        errors.append({"code": "uppercase-role", "values": uppercase_roles})

    placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(text)))
    if placeholders:
        errors.append({"code": "placeholder", "values": placeholders})

    for p_idx, paragraph in enumerate(main_body_paragraphs(doc)):
        depth = 0
        for r_idx, run in enumerate(paragraph.runs):
            if not run.text:
                continue
            name = effective_font(run, paragraph)
            size = effective_size(run, paragraph)
            italic = effective_italic(run, paragraph)
            for char in run.text:
                inside = depth > 0 or char == "("
                expected_size = 10 if inside else 12
                expected_italic = inside
                if name != FONT or size != expected_size or italic != expected_italic:
                    errors.append(
                        {
                            "code": "main-font",
                            "paragraph": p_idx,
                            "run": r_idx,
                            "text": run.text[:120],
                            "font": name,
                            "size": size,
                            "italic": italic,
                            "expected_size": expected_size,
                            "expected_italic": expected_italic,
                        }
                    )
                    break
                if char == "(":
                    depth += 1
                elif char == ")" and depth:
                    depth -= 1
        if depth:
            warnings.append({"code": "unbalanced-parentheses", "paragraph": p_idx})

    for p_idx, paragraph in enumerate(all_paragraphs(doc)):
        for r_idx, run in enumerate(paragraph.runs):
            name = effective_font(run, paragraph)
            if run.text and name not in (None, FONT):
                errors.append(
                    {
                        "code": "non-times-font",
                        "paragraph": p_idx,
                        "run": r_idx,
                        "text": run.text[:120],
                        "font": name,
                    }
                )

    if template is not None and template.exists() and sha256(path) == sha256(template):
        errors.append({"code": "unchanged-template"})

    return {
        "file": str(path.resolve()),
        "sha256": sha256(path),
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--json", type=Path, help="Optional path for the audit report")
    args = parser.parse_args()

    result = audit(args.docx, args.template)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    print(payload)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(payload + "\n", encoding="utf-8")
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
